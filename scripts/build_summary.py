#!/usr/bin/env python3
"""
Builds data/summary.json from the two NASCO scale-sample source files
(d_1SW_top15.xlsx, d_2SW_top15.xlsx).

Usage:
    python scripts/build_summary.py <path_to_1SW.xlsx> <path_to_2SW.xlsx> [output_path]

Each river/year has a candidate pool of up to 15 randomly-selected fish (per
the NASCO sampling design), of which only PER_YEAR_TARGET actually need to be
imaged for that river/year -- the rest are kept in reserve in case of poor
scale quality.

"Required"  = min(PER_YEAR_TARGET, candidate pool size) for that river/year --
              the imaging target, capped by however many candidates exist.
"Analyzed"  = however many of those already-imaged candidates count toward
              the target, i.e. min(imaged count, required).
"Excess"    = already-imaged candidates beyond the target (imaged count minus
              required, floored at 0) -- extra coverage that isn't needed to
              hit the target but exists anyway.
Required + Analyzed always describes a 0-100% target; Excess is reported
separately and is never part of that percentage.

This same logic is meant to be re-implemented by the daily refresh pipeline
(see SETUP.md) so a manual re-run of this script and the automated pipeline
always produce the identical JSON shape.
"""
import json
import sys
from collections import defaultdict
from datetime import datetime, timezone

import openpyxl

REQUIRED_COLUMNS = ["Objektnavn", "Vassdragsnr_hovedvassdrag", "Feltaar", "Bilde_skjell"]

# Per-river, per-year imaging target. The candidate pool itself is usually
# larger (up to 15) -- the remainder are reserve fish kept in case of poor
# scale quality. Confirmed with the project lead: 10 for both age classes.
PER_YEAR_TARGET = {"sw1": 10, "sw2": 10}


def load_rows(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    idx = {name: i for i, name in enumerate(header_row)}
    missing = [c for c in REQUIRED_COLUMNS if c not in idx]
    if missing:
        raise SystemExit(f"{path}: missing expected columns {missing}")

    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        river = r[idx["Objektnavn"]]
        watershed = r[idx["Vassdragsnr_hovedvassdrag"]]
        year = r[idx["Feltaar"]]
        has_image = r[idx["Bilde_skjell"]] not in (None, "")
        if river is None or year is None:
            continue
        rows.append((str(river), str(watershed) if watershed is not None else "", int(year), has_image))
    return rows


FIELDS = ["sw1Required", "sw1Analyzed", "sw1Excess", "sw2Required", "sw2Analyzed", "sw2Excess"]


def aggregate(rows_1sw, rows_2sw):
    rivers = {}  # name -> {watershedId, byYear: {year: {...}}}
    all_years = set()

    def add(rows, sw_key):
        target = PER_YEAR_TARGET[sw_key]
        # Group first: required/analyzed/excess depend on the whole
        # river+year candidate pool, not on individual rows.
        groups = defaultdict(list)  # (river, watershed, year) -> [has_image, ...]
        for river, watershed, year, has_image in rows:
            groups[(river, watershed, year)].append(has_image)

        for (river, watershed, year), images in groups.items():
            all_years.add(year)
            entry = rivers.setdefault(river, {"name": river, "watershedId": watershed, "byYear": {}})
            if not entry["watershedId"] and watershed:
                entry["watershedId"] = watershed
            yr = entry["byYear"].setdefault(str(year), {k: 0 for k in FIELDS})

            candidate_count = len(images)
            imaged_count = sum(1 for img in images if img)
            required = min(target, candidate_count)
            analyzed = min(imaged_count, required)
            excess = max(0, imaged_count - required)

            yr[f"{sw_key}Required"] += required
            yr[f"{sw_key}Analyzed"] += analyzed
            yr[f"{sw_key}Excess"] += excess

    add(rows_1sw, "sw1")
    add(rows_2sw, "sw2")

    totals = {k: 0 for k in FIELDS}
    river_list = []
    for river in rivers.values():
        river_totals = {k: 0 for k in FIELDS}
        for yr in river["byYear"].values():
            for k in river_totals:
                river_totals[k] += yr[k]
        river["totals"] = river_totals
        for k in totals:
            totals[k] += river_totals[k]
        river_list.append(river)

    river_list.sort(key=lambda r: r["name"])

    return {
        "generatedAt": datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z"),
        "years": sorted(all_years),
        "rivers": river_list,
        "totals": totals,
    }


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    path_1sw, path_2sw = sys.argv[1], sys.argv[2]
    out_path = sys.argv[3] if len(sys.argv) > 3 else "data/summary.json"

    rows_1sw = load_rows(path_1sw)
    rows_2sw = load_rows(path_2sw)
    summary = aggregate(rows_1sw, rows_2sw)

    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)

    print(f"Wrote {out_path}: {len(summary['rivers'])} rivers, years {summary['years'][0]}-{summary['years'][-1]}")
    print(f"Totals: {summary['totals']}")


if __name__ == "__main__":
    main()
