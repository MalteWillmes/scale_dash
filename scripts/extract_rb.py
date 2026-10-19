#!/usr/bin/env python3
"""
Builds data/rb_pictures.json, which build_summary.py reads, for the collaborator
RB (Radgivende Biologer). Two sources, because RB's imaging is not finished:

  MATERIAL  the project appendix workbook ("Vedlegg 1 - materiale 1989-2024").
            It has one sheet per age class (1SW, 2SW, 3SW) and one row per river
            and institution (Institusjon = NINA or RB); the year columns hold how
            many fish of material exist for that river and year (at most 10).
            For the RB rows that is what RB should image: every year it lists.

  IMAGES    the folder where the finished images are kept, one sub-folder per
            year (1999, 2000, ...), each holding the Leica exports named
            <River><Year>-<fish no>.jpg (or .tif/.png) and a "DNV <year>.xlsx"
            with one row per imaged fish. That row's "ID nr." is the image name,
            and it says which watershed the fish is from ("Vassdragsnr.") and how
            many sea winters it spent ("Sjo"), so every image is assigned to a
            river, a year and an age class without guessing from file names.

Usage:
    python scripts/extract_rb.py <Vedlegg1_materiale_1989_2024.xlsx> <RB image folder> [output_path]

Only the 1SW and 2SW sheets are used (the dashboard tracks those two age
classes): fish with 3 or more sea winters are counted and reported but not used.
NINA rows are ignored: everything the dashboard already had is institution NINA.
Other folders next to the year folders (e.g. a staging folder) are ignored, and
an image counts once however many files (jpg, tif, BandTracer png) it has.
Images that have no row in the DNV workbook cannot be assigned an age class, so
they are not counted; their number is written to `imagesScanned`.

The output has no file names or fish ids, only counts per river, age class and
year, so it is safe to publish.
"""
import json
import os
import re
import sys
from collections import Counter, defaultdict
from datetime import date

import openpyxl

AGE_SHEETS = (("1SW", 1), ("2SW", 2))
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".tif", ".tiff", ".png"}
IMAGE_ID = re.compile(r"^(?P<id>.+?\d{4}-\d+)")  # e.g. Arna2000-001 (also in Arna2000-001_BandTracer.png)

# Columns of the DNV workbooks (header row 5), looked up by name.
DNV_COLUMNS = {"id": "ID nr.", "watershed": "Vassdragsnr.", "year": "År", "sea": "Sjø", "species": "Art"}
DNV_HEADER_ROW = 5  # 1-based; data starts on the row below


def ascii_fold(text):
    # The dashboard's data writes the southern region as "Sor" (no o-slash).
    return text.replace("ø", "o").replace("Ø", "O")


def extract_material(path):
    """RB's material per river: {watershed id: {name, region, material: {sheet: {year: fish}}}}."""
    wb = openpyxl.load_workbook(path, data_only=True)
    rivers = {}
    for sheet, _ in AGE_SHEETS:
        ws = wb[sheet]
        header = [c.value for c in ws[1]]
        if "Institusjon" not in header:
            raise SystemExit(f"{sheet}: expected column 'Institusjon' not found")
        i_inst = header.index("Institusjon")
        year_at = {i: h for i, h in enumerate(header) if isinstance(h, int)}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[i_inst] != "RB":
                continue
            river_id, name = row[2].split(" ", 1)
            region = ascii_fold(row[0])
            r = rivers.setdefault(river_id, {"name": name, "region": region, "material": {}})
            if r["name"] != name or r["region"] != region:
                raise SystemExit(f"{river_id}: name/region differ between sheets ({r['name']}/{r['region']} vs {name}/{region})")
            r["material"][sheet] = {
                str(year_at[i]): int(row[i]) for i in year_at if isinstance(row[i], (int, float)) and row[i] > 0
            }
    return rivers


def find_dnv_columns(header):
    names = {str(c).strip(): i for i, c in enumerate(header) if c is not None}
    missing = [n for n in DNV_COLUMNS.values() if n not in names]
    if missing:
        raise SystemExit(f"DNV workbook is missing columns {missing}")
    return {key: names[n] for key, n in DNV_COLUMNS.items()}


def scan_images(folder, known_watersheds):
    """Counts distinct imaged fish per (watershed id, year, sea winters) from the year folders."""
    counts = Counter()
    stats = {"years": [], "images": 0, "counted": 0, "noDnvRow": 0, "otherAge": 0}
    for entry in sorted(os.listdir(folder)):
        full = os.path.join(folder, entry)
        if not (entry.isdigit() and os.path.isdir(full)):
            continue  # only year folders; anything else (e.g. a staging folder) is ignored
        year = int(entry)

        ids = set()
        for dirpath, _, files in os.walk(full):
            for f in files:
                stem, ext = os.path.splitext(f)
                m = IMAGE_ID.match(stem)
                if ext.lower() in IMAGE_EXTENSIONS and m:
                    ids.add(m.group("id"))

        dnv_path = os.path.join(full, f"DNV {year}.xlsx")
        if not os.path.exists(dnv_path):
            raise SystemExit(f"{entry}: images found but {dnv_path} is missing, so they can't be assigned to rivers")
        wb = openpyxl.load_workbook(dnv_path, data_only=True, read_only=True)
        rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
        cols = find_dnv_columns(rows[DNV_HEADER_ROW - 1])
        fish = {}
        for row in rows[DNV_HEADER_ROW:]:
            if row[cols["id"]]:
                fish[str(row[cols["id"]]).strip()] = row

        stats["years"].append(year)
        stats["images"] += len(ids)
        for image_id in sorted(ids):
            row = fish.get(image_id)
            if row is None:
                stats["noDnvRow"] += 1
                continue
            if row[cols["year"]] != year:
                raise SystemExit(f"{image_id}: DNV says year {row[cols['year']]}, but it is in the {year} folder")
            if row[cols["species"]] != "Laks":
                raise SystemExit(f"{image_id}: species {row[cols['species']]!r}, expected 'Laks'")
            watershed = str(row[cols["watershed"]]).strip()
            if watershed not in known_watersheds:
                raise SystemExit(f"{image_id}: watershed {watershed!r} is not one of RB's rivers in the appendix workbook")
            sea = row[cols["sea"]]
            if sea not in (1, 2):
                stats["otherAge"] += 1
                continue
            counts[(watershed, year, sea)] += 1
            stats["counted"] += 1
    return counts, stats


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    material_path, images_folder = sys.argv[1], sys.argv[2]
    out_path = sys.argv[3] if len(sys.argv) > 3 else "data/rb_pictures.json"

    rivers = extract_material(material_path)
    counts, stats = scan_images(images_folder, set(rivers))

    for river_id, r in rivers.items():
        r["pictures"] = {}
        for sheet, sea in AGE_SHEETS:
            by_year = {str(y): n for (w, y, s), n in sorted(counts.items()) if w == river_id and s == sea}
            if by_year:
                r["pictures"][sheet] = by_year
    rivers = dict(sorted(rivers.items()))

    doc = {
        "about": (
            "The collaborator RB (Radgivende Biologer), by river, age class and year. `material` is how many fish RB "
            "has for that river-year (appendix workbook 'Vedlegg 1 - materiale 1989-2024', Institusjon = RB rows, "
            "sheets 1SW and 2SW) and so what has to be imaged, at most 10 per river-year. `pictures` is how many fish "
            "have been imaged so far (counted from the image folders and their DNV workbooks); RB's imaging is not "
            "finished, so most years have none yet. `imagesScanned` summarises the scan. Regenerate with "
            "scripts/extract_rb.py."
        ),
        "source": os.path.basename(material_path),
        "extractedOn": date.today().isoformat(),
        "imagesScanned": {
            "years": stats["years"],
            "imageFiles": stats["images"],
            "counted": stats["counted"],
            "notCountedNoDnvRow": stats["noDnvRow"],
            "notCountedOtherAgeClass": stats["otherAge"],
        },
        "rivers": rivers,
    }
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(doc, f, ensure_ascii=False, indent=1)

    material = {s: sum(sum(r["material"].get(s, {}).values()) for r in rivers.values()) for s, _ in AGE_SHEETS}
    pictures = {s: sum(sum(r["pictures"].get(s, {}).values()) for r in rivers.values()) for s, _ in AGE_SHEETS}
    print(f"Wrote {out_path}: {len(rivers)} RB rivers")
    print(f"  material (fish to image): {material}")
    print(f"  imaged so far:            {pictures}  from {stats['images']} images in years {stats['years'][0]}-{stats['years'][-1]}")
    print(f"  not counted: {stats['noDnvRow']} without a DNV row, {stats['otherAge']} with 3+ sea winters")


if __name__ == "__main__":
    main()
